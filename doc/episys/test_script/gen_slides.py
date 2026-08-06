#!/usr/bin/env python3
# Usage:
#   pip install python-pptx    # install dependency (one-time)
#   python3 gen_slides.py      # generates overview_sl_test.pptx
#
# Output: overview_sl_test.pptx in the same directory as this script.
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color palette
DARK_BLUE = RGBColor(0x1A, 0x3A, 0x5C)
MID_BLUE = RGBColor(0x2A, 0x64, 0x96)
LIGHT_BLUE = RGBColor(0x3B, 0x82, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NEAR_BLACK = RGBColor(0x1E, 0x29, 0x3B)
GRAY = RGBColor(0x64, 0x74, 0x8B)
LIGHT_BG = RGBColor(0xF8, 0xFA, 0xFC)
TABLE_HEADER_BG = RGBColor(0x1A, 0x3A, 0x5C)
TABLE_EVEN_BG = RGBColor(0xF4, 0xF8, 0xFC)
BADGE_BLUE_BG = RGBColor(0xDB, 0xEA, 0xFE)
BADGE_BLUE_FG = RGBColor(0x1E, 0x40, 0xAF)
GREEN_BG = RGBColor(0xD1, 0xFA, 0xE5)
GREEN_FG = RGBColor(0x06, 0x5F, 0x46)
ACCENT_BLUE = RGBColor(0x2A, 0x7A, 0xE2)
BOX_BG = RGBColor(0xF0, 0xF4, 0xF8)
ORANGE_BG = RGBColor(0xFF, 0xFB, 0xEB)
ORANGE_FG = RGBColor(0x9A, 0x34, 0x12)


def add_rounded_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    return shape


def set_cell_text(cell, text, font_size=11, bold=False, color=NEAR_BLACK, alignment=PP_ALIGN.LEFT):
    cell.text = ""
    p = cell.text_frame.paragraphs[0]
    p.alignment = alignment
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE


def add_text_box(slide, left, top, width, height, text, font_size=14, color=NEAR_BLACK, bold=False, alignment=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = alignment
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.color.rgb = color
    run.font.bold = bold
    return txBox


def add_bullet_list(slide, left, top, width, height, items, font_size=13, color=NEAR_BLACK, spacing=Pt(6)):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.space_after = spacing
        p.level = 0
        if isinstance(item, tuple):
            bold_part, normal_part = item
            run_b = p.add_run()
            run_b.text = bold_part
            run_b.font.size = Pt(font_size)
            run_b.font.bold = True
            run_b.font.color.rgb = color
            run_n = p.add_run()
            run_n.text = normal_part
            run_n.font.size = Pt(font_size)
            run_n.font.color.rgb = color
        else:
            run = p.add_run()
            run.text = f"•  {item}"
            run.font.size = Pt(font_size)
            run.font.color.rgb = color
    return txBox


def add_code_block(slide, left, top, width, height, text, font_size=10):
    shape = add_rounded_rect(slide, left, top, width, height, NEAR_BLACK)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(12)
    tf.margin_right = Pt(12)
    tf.margin_top = Pt(10)
    tf.margin_bottom = Pt(10)
    lines = text.split('\n')
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
        run.font.name = "Consolas"
    return shape


def add_info_box(slide, left, top, width, height, text, bg_color=BOX_BG, accent_color=ACCENT_BLUE, font_size=12):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = accent_color
    shape.line.width = Pt(2)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(14)
    tf.margin_top = Pt(8)
    for i, line in enumerate(text.split('\n')):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = line
        run.font.size = Pt(font_size)
        run.font.color.rgb = NEAR_BLACK
    return shape


def add_table(slide, left, top, width, rows_data, col_widths, header=True, col_alignments=None, row_height=0.4):
    num_rows = len(rows_data)
    num_cols = len(rows_data[0])
    table_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, Inches(row_height * num_rows))
    table = table_shape.table
    for i, w in enumerate(col_widths):
        table.columns[i].width = w
    for r, row in enumerate(rows_data):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            align = PP_ALIGN.CENTER if (col_alignments and col_alignments[c] == "center") else PP_ALIGN.LEFT
            if r == 0 and header:
                set_cell_text(cell, val, font_size=11, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
                cell.fill.solid()
                cell.fill.fore_color.rgb = TABLE_HEADER_BG
            else:
                set_cell_text(cell, val, font_size=10, color=NEAR_BLACK, alignment=align)
                if r % 2 == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = TABLE_EVEN_BG
                else:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = WHITE
    return table_shape


def slide_header(slide, title):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.08))
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT_BLUE
    shape.line.fill.background()
    add_text_box(slide, Inches(0.6), Inches(0.25), Inches(10), Inches(0.7), title,
                 font_size=28, color=DARK_BLUE, bold=True)


# ============================================================
# SLIDE 1: Title
# ============================================================
slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # blank
bg = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = DARK_BLUE
bg.line.fill.background()

add_text_box(slide1, Inches(1.5), Inches(1.5), Inches(10), Inches(1.2),
             "OAI 5G NR Sidelink", font_size=48, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
add_text_box(slide1, Inches(1.5), Inches(2.5), Inches(10), Inches(1.0),
             "Test Framework", font_size=48, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
add_text_box(slide1, Inches(1.5), Inches(3.7), Inches(10), Inches(0.6),
             "Automated CI Test Harness for PC5 & U2N Relay Validation",
             font_size=20, color=RGBColor(0xBF, 0xDB, 0xFE), alignment=PP_ALIGN.CENTER)

badges = ["PC5 Mode 1 & 2", "CSI / PSFCH", "U2N Relay (SRAP)", "RF Sim, VRT Sim & USRP"]
# Per-badge widths sized to their text (last one is wider so "USRP" is not clipped),
# and the whole row is centered on the slide.
badge_widths = [Inches(1.5), Inches(1.2), Inches(1.7), Inches(2.2)]
badge_gap = Inches(0.2)
badge_total = sum(badge_widths) + badge_gap * (len(badges) - 1)
badge_x = int((prs.slide_width - badge_total) / 2)
for badge_text, bw in zip(badges, badge_widths):
    shape = add_rounded_rect(slide1, badge_x, Inches(4.8), bw, Inches(0.4),
                             RGBColor(0x2A, 0x64, 0x96), RGBColor(0x60, 0xA5, 0xFA))
    tf = shape.text_frame
    tf.word_wrap = False
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    run = tf.paragraphs[0].add_run()
    run.text = badge_text
    run.font.size = Pt(11)
    run.font.color.rgb = WHITE
    run.font.bold = True
    badge_x += int(bw) + int(badge_gap)

add_text_box(slide1, Inches(1.5), Inches(6.2), Inches(10), Inches(0.5),
             "OpenAirInterface — 2026", font_size=14, color=RGBColor(0x93, 0xC5, 0xFD),
             alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 2: Overview & Key Features
# ============================================================
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
slide_header(slide2, "Overview & Key Features")

add_text_box(slide2, Inches(0.6), Inches(1.1), Inches(5), Inches(0.5),
             "What It Tests", font_size=18, color=MID_BLUE, bold=True)
add_bullet_list(slide2, Inches(0.6), Inches(1.6), Inches(5.5), Inches(3.0), [
    ("Sidelink Mode 2 ", "— Direct D2D communication"),
    ("Sidelink Mode 1 ", "— Network-scheduled sidelink"),
    ("U2N Relay (SRAP) ", "— Remote UE connectivity via relay"),
    ("CSI / PSFCH ", "— Parametric feedback testing"),
    ("Uu Interface ", "— gNB <-> UE connectivity testing"),
    ("iperf3 Bandwidth Sweep ", "— UDP throughput characterization"),
    ("BLER Testing ", "— Waterfall curves across MCS 0-28 with noise sweep"),
], font_size=14)

add_text_box(slide2, Inches(7.0), Inches(1.1), Inches(5), Inches(0.5),
             "Key Features", font_size=18, color=MID_BLUE, bold=True)
add_bullet_list(slide2, Inches(7.0), Inches(1.6), Inches(5.5), Inches(3.0), [
    ("Array-based test selection ", "— simple ordered list"),
    ("Parametric sweeps ", "— CSI x PSFCH combinations"),
    ("Multi-host support ", "— 1, 2, or 3 machines via SSH"),
    ("Profile-based config ", "— pilot / regress / stress / bler"),
    ("Auto result tracking ", "— ping + PSSCH statistics"),
    ("Three radio backends ", "— rfsim (socket) / usrp (B210 HW) / vrtsim (shared-memory, local host)"),
], font_size=14)

add_info_box(slide2, Inches(0.6), Inches(5.5), Inches(12.0), Inches(0.7),
             "Two scripts:  run_sl_test.sh (test engine)  +  run_sl_test_config.sh (user configuration)",
             font_size=14)

# ============================================================
# SLIDE 3: Architecture & Execution Flow
# ============================================================
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
slide_header(slide3, "Architecture & Execution Flow")

add_text_box(slide3, Inches(0.6), Inches(1.1), Inches(5), Inches(0.5),
             "Execution Flow", font_size=18, color=MID_BLUE, bold=True)
add_code_block(slide3, Inches(0.6), Inches(1.6), Inches(5.5), Inches(4.0),
    "1. Load config (run_sl_test_config.sh)\n"
    "2. Resolve enabled_tests array\n"
    "3. For each test:\n"
    "   a. Clean up old logs\n"
    "   b. Launch UEs (local / remote SSH)\n"
    "   c. Wait for tunnel interface\n"
    "   d. Wait for PC5 sync (PSBCH decode)\n"
    "   e. Execute ping test\n"
    "   f. Collect stats (ping + PSSCH)\n"
    "   g. Kill processes & cleanup\n"
    "   h. Generate summary row\n"
    "4. Display final summary table\n"
    "\n"
    "Duration budget (ensure_ping_test_time):\n"
    "  =1: flexible, ping always runs (min 16s)\n"
    "  =0: strict, test ends after duration",
    font_size=12)

add_text_box(slide3, Inches(7.0), Inches(1.1), Inches(5), Inches(0.5),
             "File Structure", font_size=18, color=MID_BLUE, bold=True)
add_code_block(slide3, Inches(7.0), Inches(1.6), Inches(5.5), Inches(4.0),
    "run_sl_test.sh            # Test engine\n"
    "run_sl_test_config.sh     # Configuration\n"
    "~/.ssh/config             # SSH host aliases\n"
    "~/openairinterface5g/     # OAI source\n"
    "\n"
    "<base_dir>/\n"
    "  latest -> test_<ts>/    # Symlink\n"
    "  test_<timestamp>/       # Results\n"
    "    test_summary_*.csv\n"
    "    iperf3_summary_*.csv / *.png\n"
    "    result_*_<test>_<ts>.log\n"
    "    ping_result_*.txt",
    font_size=12)

add_info_box(slide3, Inches(7.0), Inches(5.8), Inches(5.5), Inches(0.6),
             "base_dir priority:  -d flag > base_log_dir in config > script directory",
             font_size=12)

# ============================================================
# SLIDE 4: Multi-Host Testing & SSH Setup
# ============================================================
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
slide_header(slide4, "Multi-Host Testing")

add_text_box(slide4, Inches(0.6), Inches(1.1), Inches(5), Inches(0.5),
             "Test Topologies", font_size=18, color=MID_BLUE, bold=True)
add_code_block(slide4, Inches(0.6), Inches(1.6), Inches(5.5), Inches(4.5),
    "Single-Host (Mode 2):\n"
    "  SyncRef UE <--rfsim--> Nearby UE\n"
    "  (both on localhost)\n"
    "\n"
    "Two-Host (Mode 2):\n"
    "  Local:  SyncRef UE (rfsim server)\n"
    "  Remote: Nearby UE  (rfsim client)\n"
    "\n"
    "Three-Host (Mode 1 SRAP):\n"
    "  gNB host:       gNB + 5G Core\n"
    "  relay_ue host:  Relay UE (Uu + PC5)\n"
    "  remote_ue host: Remote UE (PC5)",
    font_size=12)

add_text_box(slide4, Inches(7.0), Inches(1.1), Inches(5), Inches(0.5),
             "Host Assignment per Test Mode", font_size=18, color=MID_BLUE, bold=True)

host_data = [
    ["SL Mode", "Type", "Hosts", "gNB", "SyncRef UE", "Nearby UE"],
    ["Mode 2", "RFSIM", "1", "—", "local", "local"],
    ["Mode 2", "RFSIM", "2", "—", "local", "remote_ue"],
    ["Mode 2", "USRP", "2", "—", "local", "remote_ue"],
    ["Mode 2", "VRTSIM", "1", "—", "local", "local"],
    ["Mode 1 (SRAP)", "RFSIM", "1", "local", "local", "local"],
    ["Mode 1 (SRAP)", "RFSIM", "3", "gNB, local", "relay_ue", "remote_ue"],
    ["Mode 1 (SRAP)", "USRP", "3", "local", "relay_ue", "remote_ue"],
    ["Mode 1 (SRAP)", "VRTSIM", "1", "local", "local", "local"],
    ["Uu", "RFSIM", "1", "local", "—", "local (nrUE)"],
    ["Uu", "RFSIM", "2", "local", "—", "nr_ue (nrUE)"],
    ["Uu", "USRP", "2", "local", "—", "nr_ue (nrUE)"],
    ["Uu", "VRTSIM", "1", "local", "—", "local (nrUE)"],
]
add_table(slide4, Inches(7.0), Inches(1.7), Inches(5.5), host_data,
          [Inches(1.0), Inches(0.6), Inches(0.5), Inches(0.8), Inches(1.1), Inches(1.5)],
          col_alignments=["left", "center", "center", "center", "center", "center"])

add_info_box(slide4, Inches(7.0), Inches(6.0), Inches(5.5), Inches(0.7),
             "SSH aliases defined in ~/.ssh/config\nPasswordless SSH required: ssh-copy-id <alias>",
             bg_color=ORANGE_BG, accent_color=RGBColor(0xF5, 0x9E, 0x0B), font_size=12)

# ============================================================
# SLIDE 5: Test Selection & Configuration
# ============================================================
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
slide_header(slide5, "Test Selection & Configuration")

add_text_box(slide5, Inches(0.6), Inches(1.1), Inches(5), Inches(0.5),
             "Test Profiles", font_size=18, color=MID_BLUE, bold=True)

profile_data = [
    ["Profile", "Repeats", "MCS", "Duration", "SNR/Atten", "Gain", "LDPC Iter", "Use Case"],
    ["pilot", "1", "1", "30s", "0 / 20dB", "20/110", "30", "Quick smoke test"],
    ["regress", "1", "1, 9", "30s", "0 / 20dB", "20/110", "30", "Regression validation"],
    ["stress", "3", "9, 16, 28", "300s", "0 / 20-60dB", "20/110", "30", "Long-term stability"],
    ["bler", "12", "0-28", "85s", "Noise sweep", "N/A", "30", "BLER waterfall curves"],
]
add_table(slide5, Inches(0.6), Inches(1.6), Inches(5.5), profile_data,
          [Inches(0.7), Inches(0.55), Inches(0.7), Inches(0.55), Inches(0.8), Inches(0.5), Inches(0.55), Inches(1.2)],
          col_alignments=["left", "center", "center", "center", "center", "center", "center", "left"])

add_text_box(slide5, Inches(0.6), Inches(3.8), Inches(5), Inches(0.5),
             "Per-Profile Test Lists", font_size=18, color=MID_BLUE, bold=True)
add_bullet_list(slide5, Inches(0.6), Inches(4.3), Inches(5.5), Inches(2.5), [
    "pilot_tests / regress_tests / stress_tests",
    "slmode2_basic_tests — entire group",
    "slmode2_basic_tests[0:2] — range (indices 0-2)",
    "slmode2_basic_tests[0] — single test",
    "test_name:csi:psfch — parametric combo",
], font_size=13)

add_text_box(slide5, Inches(7.0), Inches(1.1), Inches(5), Inches(0.5),
             "Additional Config Settings", font_size=18, color=MID_BLUE, bold=True)
add_code_block(slide5, Inches(7.0), Inches(1.6), Inches(5.5), Inches(2.8),
    "# Log output directory (default: script dir)\n"
    "base_log_dir=\"~/openairinterface5g\"\n"
    "\n"
    "# Use external clock source for USRP (0=internal, 1=external) \n"
    "use_external_clock=1\n"
    "# Use standalone mode (0=disabled, 1=enabled)\n"
    "use_sa=0\n"
    "# Use gnome-terminal (0=bash, 1=gnome)\n"
    "use_gnome=1\n"
    "# Ensure ping test time (0=strict, 1=flexible)\n"
    "ensure_ping_test_time=1\n"
    "\n"
    "# USRP serial numbers (SL Mode 1 relay only)\n"
    "RELAY_UE_USRP_SN_FOR_UU=340EA03\n"
    "RELAY_UE_USRP_SN_FOR_SL=340EA3B",
    font_size=9)

add_text_box(slide5, Inches(7.0), Inches(4.3), Inches(5), Inches(0.5),
             "Running Tests", font_size=18, color=MID_BLUE, bold=True)
add_code_block(slide5, Inches(7.0), Inches(4.8), Inches(5.5), Inches(1.5),
    "# use config settings (or defaults)\n"
    "./run_sl_test.sh\n"
    "# override log dir and gnome terminal\n"
    "./run_sl_test.sh -d ~/openairinterface5g\n"
    "./run_sl_test.sh -g 1\n"
    "./run_sl_test.sh -d ~/openairinterface5g -g 1",
    font_size=11)

add_info_box(slide5, Inches(7.0), Inches(6.5), Inches(5.5), Inches(0.6),
             "Priority:  CLI flags (-d, -g) > config file > defaults",
             font_size=12)

# ============================================================
# SLIDE 5.5: Four-Tier Configuration System
# ============================================================
slide5_5 = prs.slides.add_slide(prs.slide_layouts[6])
slide_header(slide5_5, "Four-Tier Configuration System")

add_text_box(slide5_5, Inches(0.6), Inches(1.0), Inches(12), Inches(0.4),
             "Override MCS and duration at four levels:  test-specific > slice-specific > group-specific > profile-default",
             font_size=14, color=MID_BLUE)

add_text_box(slide5_5, Inches(0.6), Inches(1.6), Inches(5), Inches(0.5),
             "Configuration Examples", font_size=18, color=MID_BLUE, bold=True)
add_code_block(slide5_5, Inches(0.6), Inches(2.1), Inches(5.5), Inches(3.8),
    "# In run_sl_test_config.sh\n"
    "\n"
    "# Group-level (all tests in group)\n"
    "group_specific_duration[\"slmode2_basic_tests\"]=30\n"
    "group_specific_mcs[\"slmode2_basic_tests\"]=\"16,28\"\n"
    "\n"
    "# Slice-level (specific indices)\n"
    "# [0:2] range, [0,2] pick-list, [2] single\n"
    "group_specific_duration[\"slmode2_basic_tests[2]\"]=60\n"
    "group_specific_mcs[\"slmode2_basic_tests[2]\"]=\"12,16,20,24,28\"\n"
    "\n"
    "# Test-specific (individual test)\n"
    "test_specific_duration[\"usrp_B210_pc5_ping_test\"]=90\n"
    "test_specific_mcs[\"usrp_B210_pc5_ping_test\"]=\"20,24,28\"",
    font_size=11)

add_text_box(slide5_5, Inches(7.0), Inches(1.6), Inches(5), Inches(0.5),
             "Key Features", font_size=18, color=MID_BLUE, bold=True)
add_bullet_list(slide5_5, Inches(7.0), Inches(2.1), Inches(5.5), Inches(2.0), [
    ("MCS formats: ", "Comma \"16,28\", space \"16 28\", or seq \"$(seq 0 1 10)\""),
    ("Auto-discovery: ", "Groups ending with _basic_tests, _iperf3_tests, _csi_psfch_tests"),
    ("Excluded: ", "pilot_tests, regress_tests, stress_tests (profile arrays)"),
], font_size=13, spacing=Pt(10))

add_info_box(slide5_5, Inches(7.0), Inches(4.5), Inches(5.5), Inches(1.2),
             "Priority order:\n"
             "1. Test-specific → 2. Slice-specific → 3. Group-specific → 4. Profile-default",
             bg_color=GREEN_BG, accent_color=RGBColor(0x10, 0xB9, 0x81), font_size=14)

# ============================================================
# SLIDE 6: Predefined Test Groups
# ============================================================
slide6 = prs.slides.add_slide(prs.slide_layouts[6])
slide_header(slide6, "Predefined Test Groups")

groups_data = [
    ["Group", "Idx", "Test Case"],
    ["uu_basic_tests", "0", "rfsim_uu_ping_test_on_local_host"],
    ["", "1", "vrtsim_uu_ping_test_on_local_host"],
    ["", "2", "rfsim_uu_ping_test_on_two_hosts"],
    ["", "3", "usrp_B210_uu_ping_test_on_two_hosts"],
    ["slmode2_basic_tests", "0", "rfsim_pc5_ping_test_on_local_host"],
    ["", "1", "vrtsim_pc5_ping_test_on_local_host"],
    ["", "2", "rfsim_pc5_ping_test_on_two_hosts"],
    ["", "3", "usrp_B210_pc5_ping_test_on_two_hosts"],
    ["slmode2_csi_psfch_tests", "0", "rfsim_pc5_csi_acquisition_psfch_period_test_on_local_host"],
    ["", "1", "vrtsim_pc5_csi_acquisition_psfch_period_test_on_local_host"],
    ["", "2", "rfsim_pc5_csi_acquisition_psfch_period_test_on_two_hosts"],
    ["", "3", "usrp_B210_pc5_csi_acquisition_psfch_period_test_on_two_hosts"],
    ["slmode2_iperf3_tests", "0", "rfsim_pc5_iperf3_test_on_local_host"],
    ["", "1", "vrtsim_pc5_iperf3_test_on_local_host"],
    ["", "2", "rfsim_pc5_iperf3_test_on_two_hosts"],
    ["", "3", "usrp_B210_pc5_iperf3_test_on_two_hosts"],
    ["slmode1_basic_tests", "0", "rfsim_slmode1_srap_ping_test_on_local_host"],
    ["", "1", "vrtsim_slmode1_srap_ping_test_on_local_host"],
    ["", "2", "rfsim_slmode1_srap_ping_test_on_three_hosts"],
    ["", "3", "usrp_B210_slmode1_srap_ping_test_on_three_hosts"],
    ["slmode1_iperf3_tests", "0", "rfsim_slmode1_srap_iperf3_test_on_local_host"],
    ["", "1", "vrtsim_slmode1_srap_iperf3_test_on_local_host"],
    ["", "2", "rfsim_slmode1_srap_iperf3_test_on_three_hosts"],
    ["", "3", "usrp_B210_slmode1_srap_iperf3_test_on_three_hosts"],
    ["slmode1_csi_psfch_tests", "0", "rfsim_slmode1_srap_csi_acquisition_psfch_period_test_on_local_host"],
    ["", "1", "vrtsim_slmode1_srap_csi_acquisition_psfch_period_test_on_local_host"],
    ["", "2", "rfsim_slmode1_srap_csi_acquisition_psfch_period_test_on_three_hosts"],
    ["", "3", "usrp_B210_slmode1_srap_csi_acquisition_psfch_period_test_on_three_hosts"],
]
add_table(slide6, Inches(0.6), Inches(0.85), Inches(6.5), groups_data,
          [Inches(2.0), Inches(0.5), Inches(4.0)],
          col_alignments=["left", "center", "left"], row_height=0.12)

add_text_box(slide6, Inches(7.8), Inches(1.1), Inches(5), Inches(0.5),
             "Custom Groups", font_size=18, color=MID_BLUE, bold=True)
add_code_block(slide6, Inches(7.8), Inches(1.6), Inches(4.8), Inches(1.2),
    "regress_tests=(\n"
    "    slmode2_basic_tests\n"
    "    slmode2_csi_psfch_tests\n"
    ")",
    font_size=11)

add_text_box(slide6, Inches(7.8), Inches(3.0), Inches(5), Inches(0.5),
             "CSI/PSFCH Parameter Format", font_size=18, color=MID_BLUE, bold=True)
add_bullet_list(slide6, Inches(7.8), Inches(3.5), Inches(4.8), Inches(2.0), [
    "test_name — all 8 combos (CSI: 0,1 x PSFCH: 0,1,2,3)",
    "test_name:0:1 — CSI=0, PSFCH=1 only",
    "test_name:1: — CSI=1, all PSFCH values",
    "test_name::2 — PSFCH=2, all CSI values",
], font_size=12, spacing=Pt(8))

add_info_box(slide6, Inches(7.8), Inches(5.5), Inches(4.8), Inches(0.7),
             "CSI: sl_CSI_Acquisition (0=enabled, 1=disabled)\nPSFCH: sl_PSFCH_Period (0, 1, 2, 3)",
             font_size=12)

add_info_box(slide6, Inches(7.8), Inches(6.4), Inches(4.8), Inches(0.5),
             "Groups resolve recursively — a group can contain other group names",
             bg_color=GREEN_BG, accent_color=RGBColor(0x10, 0xB9, 0x81), font_size=12)

# ============================================================
# SLIDE 7: Available Test Cases
# ============================================================
slide7 = prs.slides.add_slide(prs.slide_layouts[6])
slide_header(slide7, "Available Test Cases")

test_data = [
    ["Test Case", "Mode", "Hosts", "Description"],
    ["rfsim_uu_ping_test_on_local_host", "Uu", "1", "gNB <-> UE ping on localhost (requires 5G Core)"],
    ["rfsim_uu_ping_test_on_two_hosts", "Uu", "2", "gNB <-> UE ping across two machines (requires 5G Core)"],
    ["rfsim_pc5_ping_test_on_local_host", "SL Mode 2", "1", "Basic PC5 ping: SyncRef <-> Nearby UE"],
    ["rfsim_pc5_ping_test_on_two_hosts", "SL Mode 2", "2", "PC5 ping across two machines"],
    ["rfsim_pc5_csi_acquisition_psfch_*", "SL Mode 2", "1-2", "CSI/PSFCH parametric sweep (up to 8 combos)"],
    ["rfsim_pc5_iperf3_test_on_*", "SL Mode 2", "1-2", "iperf3 bandwidth sweep over PC5"],
    ["rfsim_slmode1_srap_ping_*", "SL Mode 1", "1, 3", "U2N relay ping via SRAP"],
    ["rfsim_slmode1_srap_iperf3_*", "SL Mode 1", "1, 3", "U2N relay iperf3 bandwidth sweep"],
    ["rfsim_slmode1_srap_csi_acquisition_psfch_*", "SL Mode 1", "1, 3", "U2N relay CSI/PSFCH parametric sweep (up to 8 combos)"],
    ["usrp_B210_*", "All", "2-3", "USRP hardware variants of above tests"],
    ["vrtsim_uu_ping_test_on_local_host", "Uu", "1", "gNB <-> UE ping over shared-memory radio"],
    ["vrtsim_pc5_ping_test_on_local_host", "SL Mode 2", "1", "PC5 ping over shared-memory radio"],
    ["vrtsim_pc5_csi_acquisition_psfch_*", "SL Mode 2", "1", "PC5 CSI/PSFCH parametric sweep over shared-memory radio (8 combos)"],
    ["vrtsim_pc5_iperf3_test_on_local_host", "SL Mode 2", "1", "PC5 iperf3 bandwidth sweep over shared-memory radio"],
    ["vrtsim_slmode1_srap_ping_test_on_local_host", "SL Mode 1", "1", "U2N relay ping (SRAP) over shared-memory radio; launch-order/timing sensitive"],
    ["vrtsim_slmode1_srap_csi_acquisition_psfch_*", "SL Mode 1", "1", "U2N relay CSI/PSFCH parametric sweep over shared-memory radio (8 combos)"],
    ["vrtsim_slmode1_srap_iperf3_test_on_local_host", "SL Mode 1", "1", "U2N relay iperf3 sweep over shared-memory radio; client binds oaitun_ue2"],
]
add_table(slide7, Inches(0.6), Inches(0.85), Inches(12.0), test_data,
          [Inches(4.5), Inches(1.5), Inches(1.0), Inches(5.0)],
          col_alignments=["left", "center", "center", "left"], row_height=0.22)

add_info_box(slide7, Inches(0.6), Inches(5.3), Inches(12.0), Inches(0.8),
             "Pass criteria: >= 60% ICMP ping success rate   |   PSSCH stats: TX/RX counters extracted from UE logs for each direction",
             bg_color=GREEN_BG, accent_color=RGBColor(0x10, 0xB9, 0x81), font_size=14)

# ============================================================
# SLIDE 8: Test Results & PSSCH Statistics
# ============================================================
slide8 = prs.slides.add_slide(prs.slide_layouts[6])
slide_header(slide8, "Test Results & PSSCH Statistics")

add_text_box(slide8, Inches(0.6), Inches(1.1), Inches(10), Inches(0.5),
             "Summary Table Output", font_size=18, color=MID_BLUE, bold=True)
add_code_block(slide8, Inches(0.6), Inches(1.6), Inches(12.0), Inches(1.6),
    "Test Name                                      | Itrn | Hosts | MCS | Runtime | Ping  | PSSCH Rate1     | PSSCH Rate2     | Total | Result\n"
    "==========================================================================================================================================\n"
    "rfsim_pc5_ping_test_on_local_host              |    1 |     1 |   9 |     40s |  100% | 33/33 (100%)    | 33/33 (100%)    |  100% |   PASS\n"
    "rfsim_pc5_csi_psfch_*_csi0_psfch1              |    1 |     2 |   9 |     44s |  100% | 12/12 (100%)    | 8/8 (100%)      |  100% |   PASS\n"
    "rfsim_slmode1_srap_ping_test_on_local_host     |    1 |     1 |   9 |     58s |  100% | 36/36 (100%)    | 38/38 (100%)    |  100% |   PASS",
    font_size=10)

add_text_box(slide8, Inches(0.6), Inches(3.6), Inches(5), Inches(0.5),
             "PSSCH Rate Calculation", font_size=18, color=MID_BLUE, bold=True)
add_bullet_list(slide8, Inches(0.6), Inches(4.1), Inches(5.5), Inches(2.5), [
    ("Rate1: ", "SyncRef TX -> Nearby RX  (RX_nearby / TX_syncref)"),
    ("Rate2: ", "Nearby TX -> SyncRef RX  (RX_syncref / TX_nearby)"),
    ("Total: ", "Aggregated both directions\n       (RX_sync + RX_nearby) / (TX_sync + TX_nearby)"),
], font_size=13, spacing=Pt(10))

add_text_box(slide8, Inches(7.0), Inches(3.6), Inches(5), Inches(0.5),
             "Log Files (per test)", font_size=18, color=MID_BLUE, bold=True)
add_bullet_list(slide8, Inches(7.0), Inches(4.1), Inches(5.5), Inches(2.5), [
    "test_summary_*.csv — summary table",
    "iperf3_summary_*.csv / *.png — iperf3 bandwidth/loss plot",
    "result_<component>_<test_name>_<ts>.log — softmodem output",
    "ping_result_*.txt — ping output",
    "iperf3_server_*.txt / iperf3_client_*.txt — iperf3 logs",
], font_size=12, spacing=Pt(4))

add_info_box(slide8, Inches(7.0), Inches(6.0), Inches(5.5), Inches(0.8),
             "Remote logs captured via tee, then moved to log folder.\nLog file list defined in softmodem_log_files in config.", font_size=12)

# ============================================================
# SLIDE 9: Extending the Framework
# ============================================================
slide9 = prs.slides.add_slide(prs.slide_layouts[6])
slide_header(slide9, "Extending the Framework")

add_text_box(slide9, Inches(0.6), Inches(1.1), Inches(5), Inches(0.5),
             "1. Define Test Function", font_size=18, color=MID_BLUE, bold=True)
add_code_block(slide9, Inches(0.6), Inches(1.6), Inches(5.5), Inches(4.8),
    "my_custom_test() {\n"
    "    echo \"==== Testing ${FUNCNAME[0]} ====\"\n"
    "    [[ $# -ge 1 ]] && duration=$1\n"
    "    [[ $# -ge 2 ]] && test_type=$2\n"
    "    [[ $# -ge 3 ]] && mcs=$3\n"
    "    [[ $# -ge 4 ]] && iteration=$4\n"
    "\n"
    "    local start_time=$(date +%s) sl_mode=2\n"
    "\n"
    "    run_syncref_cmd $test_type $mcs $sl_mode \"local\"\n"
    "    run_nearby_cmd  $test_type $mcs $sl_mode \"local\"\n"
    "    sleep 5\n"
    "\n"
    "    evaluate_ping_test \"local\" \"oaitun_ue1\" \\\n"
    "        \"10.0.0.100\" $sl_mode \"${FUNCNAME[0]}\"\n"
    "\n"
    "    local end_time=$(date +%s)\n"
    "    print_runtime $start_time $end_time\n"
    "    print_test_summary \"${FUNCNAME[0]}\" \\\n"
    "        \"$iteration\" \"1\" \"$mcs\" ...\n"
    "}",
    font_size=10)

add_text_box(slide9, Inches(7.0), Inches(1.1), Inches(5), Inches(0.5),
             "2. Add rfsim/usrp Wrappers", font_size=18, color=MID_BLUE, bold=True)
add_code_block(slide9, Inches(7.0), Inches(1.6), Inches(5.5), Inches(2.0),
    "rfsim_my_custom_test_on_local_host() {\n"
    "    local test_type=\"rfsim\"\n"
    "    my_custom_test $duration $test_type $mcs $iteration\n"
    "}\n"
    "\n"
    "usrp_B210_my_custom_test_on_local_host() {\n"
    "    local test_type=\"usrp\"\n"
    "    my_custom_test $duration $test_type $mcs $iteration\n"
    "}",
    font_size=10)

add_text_box(slide9, Inches(7.0), Inches(3.9), Inches(5), Inches(0.5),
             "3. Enable in Config", font_size=18, color=MID_BLUE, bold=True)
add_code_block(slide9, Inches(7.0), Inches(4.4), Inches(5.5), Inches(0.8),
    "enabled_tests=(\n"
    "    rfsim_my_custom_test_on_local_host\n"
    ")",
    font_size=10)

add_text_box(slide9, Inches(7.0), Inches(5.4), Inches(5), Inches(0.5),
             "4. Custom Profiles", font_size=18, color=MID_BLUE, bold=True)
add_code_block(slide9, Inches(7.0), Inches(5.9), Inches(5.5), Inches(1.4),
    "custom_tests=( slmode2_basic_tests )\n"
    "elif [[ $test_profile == \"custom\" ]]; then\n"
    "    enabled_tests=(\"${custom_tests[@]}\")\n"
    "    num_repeat=2; mcs_array=(1 5 9)\n"
    "    duration=60; max_ldpc_iterations=30\n"
    "    atten_array=(20 30 40)",
    font_size=10)

# ============================================================
# SLIDE 10: Troubleshooting & References
# ============================================================
slide10 = prs.slides.add_slide(prs.slide_layouts[6])
slide_header(slide10, "Troubleshooting & References")

add_text_box(slide10, Inches(0.6), Inches(1.1), Inches(6), Inches(0.5),
             "Common Issues", font_size=18, color=MID_BLUE, bold=True)

trouble_data = [
    ["Problem", "Solution"],
    ["\"No such device\" for oaitun_ue1", "UE not initialized — increase sleep after launch"],
    ["Ping fails (<60%)", "Check ip addr show oaitun_ue1; increase duration"],
    ["PSSCH stats show 0/0 or N/A", "Check sidelink logs; verify remote logs copied"],
    ["SSH command fails", "ssh remote_ue hostname — verify passwordless auth"],
    ["Remote UE config not updated", "Check REMOTE_USER var; verify paths on remote"],
    ["SRAP relay test fails", "docker ps | grep oai-amf — verify 5G Core"],
    ["USRP attenuator error", "curl http://169.254.10.10/ — check reachability"],
    ["Config changes don't persist", "Script restores defaults before each test"],
]
add_table(slide10, Inches(0.6), Inches(1.5), Inches(6.5), trouble_data,
          [Inches(2.5), Inches(4.0)])

add_text_box(slide10, Inches(7.8), Inches(1.1), Inches(5), Inches(0.5),
             "Prerequisites", font_size=18, color=MID_BLUE, bold=True)
add_bullet_list(slide10, Inches(7.8), Inches(1.6), Inches(4.8), Inches(2.0), [
    "nr-uesoftmodem, nr-softmodem, nr-cuup",
    "gnome-terminal for parallel process launch",
    "Docker-based OAI 5G Core (SRAP/Uu tests)",
    "USRP B210 + RF attenuator (hardware tests)",
    "Passwordless SSH to all remote hosts",
], font_size=12, spacing=Pt(6))

add_text_box(slide10, Inches(7.8), Inches(3.8), Inches(5), Inches(0.5),
             "References", font_size=18, color=MID_BLUE, bold=True)
add_bullet_list(slide10, Inches(7.8), Inches(4.3), Inches(4.8), Inches(2.0), [
    ("Sidelink docs: ", "~/openairinterface5g/doc/episys/README_SL.md"),
    ("Config files: ", "~/openairinterface5g/targets/.../NR-SIDELINK/CONF/"),
    ("3GPP specs: ", "TS 38.331, TS 38.321, TS 38.211"),
    ("RF simulator: ", "~/openairinterface5g/radio/rfsimulator/"),
], font_size=12, spacing=Pt(6))

add_info_box(slide10, Inches(7.8), Inches(6.2), Inches(4.8), Inches(0.5),
             "License: OAI Public License V1.1", font_size=12)

# Save (in the same directory as this script, so the deck can be checked in with the repo)
output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "overview_sl_test.pptx")
prs.save(output_path)
print(f"Saved to {output_path}")
